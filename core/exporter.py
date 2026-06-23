"""
CNKI Desktop - 导出模块
支持导出为 TXT / XLSX / DOCX 格式
"""
import os
from typing import List

from core.citation import format_citation, format_citations_batch


def export_titles_txt(papers: list, path: str):
    """导出论文标题列表为TXT文件"""
    with open(path, "w", encoding="utf-8") as f:
        for i, p in enumerate(papers, 1):
            title = getattr(p, 'title', '') if not isinstance(p, dict) else p.get('title', '')
            f.write(f"{i}. {title}\n")


def export_titles_xlsx(papers: list, path: str):
    """导出论文信息为Excel文件"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "论文列表"

    # 表头样式
    header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="134E5E", end_color="134E5E", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # 表头
    headers = ["序号", "篇名", "作者", "刊名/来源", "年份", "DOI", "被引", "下载"]
    ws.append(headers)

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # 数据
    cell_font = Font(name="微软雅黑", size=10)
    cell_align = Alignment(vertical="center", wrap_text=True)

    for i, p in enumerate(papers, 1):
        def g(attr, default=""):
            val = getattr(p, attr, None) if not isinstance(p, dict) else p.get(attr)
            if isinstance(val, list):
                return ", ".join(str(v) for v in val) if val else default
            return str(val) if val else default

        row_data = [
            i,
            g('title'),
            g('authors'),
            g('journal'),
            g('year'),
            g('doi'),
            g('cited_count'),
            g('download_count'),
        ]
        ws.append(row_data)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=i + 1, column=col_idx)
            cell.font = cell_font
            cell.alignment = cell_align
            cell.border = thin_border

    # 列宽
    col_widths = [6, 60, 30, 20, 8, 25, 8, 8]
    for idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + idx) if idx <= 26 else 'A'].width = width

    wb.save(path)


def export_citations_txt(papers: list, path: str):
    """导出GB/T 7714-2025格式引文为TXT文件"""
    with open(path, "w", encoding="utf-8") as f:
        for i, p in enumerate(papers, 1):
            f.write(format_citation(p, i) + "\n\n")


def export_citations_docx(papers: list, path: str):
    """导出GB/T 7714-2025格式引文为Word文件"""
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # 标题
    heading = doc.add_heading("参考文献列表", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 副标题
    subtitle = doc.add_paragraph("GB/T 7714-2025 格式")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.font.size = Pt(10)
        run.font.color.rgb = None

    doc.add_paragraph("")  # 空行

    # 引文列表
    for i, p in enumerate(papers, 1):
        citation_text = format_citation(p, i)
        para = doc.add_paragraph(citation_text)
        para.paragraph_format.space_after = Pt(6)
        para.paragraph_format.line_spacing = 1.5
        for run in para.runs:
            run.font.name = "宋体"
            run.font.size = Pt(10.5)  # 五号

    doc.save(path)


def export_abstracts_txt(papers: list, path: str):
    """导出论文摘要为TXT文件"""
    with open(path, "w", encoding="utf-8") as f:
        for i, p in enumerate(papers, 1):
            title = getattr(p, 'title', '')
            authors = getattr(p, 'authors', [])
            abstract = getattr(p, 'abstract', '')
            abstract_en = getattr(p, 'abstract_en', '')

            if isinstance(authors, list):
                authors_str = ", ".join(authors)
            else:
                authors_str = str(authors)

            f.write(f"[{i}] {title}\n")
            f.write(f"作者: {authors_str}\n")
            f.write(f"摘要: {abstract}\n")
            if abstract_en:
                f.write(f"Abstract: {abstract_en}\n")
            f.write("\n" + "-" * 60 + "\n\n")


def export_abstracts_docx(papers: list, path: str):
    """导出论文摘要为Word文件"""
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    heading = doc.add_heading("论文摘要汇编", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    for i, p in enumerate(papers, 1):
        title = getattr(p, 'title', '')
        authors = getattr(p, 'authors', [])
        abstract = getattr(p, 'abstract', '')
        abstract_en = getattr(p, 'abstract_en', '')

        if isinstance(authors, list):
            authors_str = ", ".join(authors)
        else:
            authors_str = str(authors)

        doc.add_heading(f"[{i}] {title}", level=2)

        para_authors = doc.add_paragraph(f"作者: {authors_str}")
        for run in para_authors.runs:
            run.font.size = Pt(10.5)

        if abstract:
            para_ab = doc.add_paragraph(f"摘要: {abstract}")
            para_ab.paragraph_format.line_spacing = 1.5
            for run in para_ab.runs:
                run.font.size = Pt(10.5)

        if abstract_en:
            para_en = doc.add_paragraph(f"Abstract: {abstract_en}")
            para_en.paragraph_format.line_spacing = 1.5
            for run in para_en.runs:
                run.font.size = Pt(10.5)

        doc.add_page_break()

    doc.save(path)


def export_all(papers: list, output_dir: str, options: dict,
               progress_callback=None):
    """统一导出接口。

    Args:
        papers: Paper对象列表
        output_dir: 输出目录
        options: 导出选项 dict:
            - titles: bool, titles_format: "txt"|"xlsx"
            - citations: bool, citations_format: "txt"|"docx"
            - abstracts: bool, abstracts_format: "txt"|"docx"
        progress_callback: 可选的进度回调函数 (current, total, message)
    """
    os.makedirs(output_dir, exist_ok=True)
    total_steps = sum(1 for k in ['titles', 'citations', 'abstracts'] if options.get(k))
    current = 0

    results = {}

    if options.get('titles'):
        current += 1
        if progress_callback:
            progress_callback(current, total_steps, "导出论文标题...")
        fmt = options.get('titles_format', 'txt')
        if fmt == 'xlsx':
            path = os.path.join(output_dir, "论文列表.xlsx")
            export_titles_xlsx(papers, path)
        else:
            path = os.path.join(output_dir, "论文列表.txt")
            export_titles_txt(papers, path)
        results['titles'] = path

    if options.get('citations'):
        current += 1
        if progress_callback:
            progress_callback(current, total_steps, "导出引文...")
        fmt = options.get('citations_format', 'txt')
        if fmt == 'docx':
            path = os.path.join(output_dir, "参考文献_GB-T_7714.docx")
            export_citations_docx(papers, path)
        else:
            path = os.path.join(output_dir, "参考文献_GB-T_7714.txt")
            export_citations_txt(papers, path)
        results['citations'] = path

    if options.get('abstracts'):
        current += 1
        if progress_callback:
            progress_callback(current, total_steps, "导出摘要...")
        fmt = options.get('abstracts_format', 'txt')
        if fmt == 'docx':
            path = os.path.join(output_dir, "论文摘要汇编.docx")
            export_abstracts_docx(papers, path)
        else:
            path = os.path.join(output_dir, "论文摘要汇编.txt")
            export_abstracts_txt(papers, path)
        results['abstracts'] = path

    return results
